"""
Per-join-method implementations for LABES enrichment.

Each indicator's `join_method` field in config/labes_indicators.yaml maps to a
function here. The orchestrator (`labes_enrich.py`) pre-loads a `JoinContext`
once and passes it to every join — that way the expensive point-in-polygon
lookups for the manifest (municipality, biogeoregion) happen exactly once and
all 30+ downstream joins reuse the result.
"""

from __future__ import annotations

import json
import logging
import re
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Mapping, Optional

import geopandas as gpd
import numpy as np
import pandas as pd
import rasterio
from rasterio.warp import transform as rio_transform
from shapely import affinity

logger = logging.getLogger("labes_joins")

TARGET_CRS = "EPSG:2056"
PROJECT_ROOT = Path("/home/jubooz/landscape_signatures")

BIOGEOREGION_NAMES = {
    1: "jura",
    2: "mittelland",
    3: "alpennordflanke",
    4: "westliche_zentralalpen",
    5: "oestliche_zentralalpen",
    6: "alpensuedflanke",
}


@dataclass
class JoinContext:
    """Pre-loaded spatial lookups that every indicator may use."""

    manifest: pd.DataFrame  # global_index, lv95_easting, lv95_northing
    biogeoregion_keys: dict[str, int] = field(default_factory=dict)
    municipality_bfs_per_point: Optional[pd.Series] = None  # indexed by global_index
    biogeoregion_per_point: Optional[pd.Series] = None      # indexed by global_index, codes 1..6


# -------------------------------------------------------------------
# Context bootstrap
# -------------------------------------------------------------------


def _read_layer(path: Path, layer: Optional[str] = None) -> gpd.GeoDataFrame:
    kwargs = {}
    if layer:
        kwargs["layer"] = layer
    gdf = gpd.read_file(path, **kwargs)
    if gdf.crs is None:
        raise ValueError(f"{path} has no CRS")
    try:
        projected = gdf.to_crs(TARGET_CRS)
        if np.isfinite(projected.total_bounds).all():
            return projected
    except Exception:
        pass
    if gdf.crs.to_epsg() == 21781:
        logger.warning("[%s] LV03 fallback +2_000_000 / +1_000_000", path)
        gdf = gdf.copy()
        gdf.geometry = gdf.geometry.apply(
            lambda geom: affinity.translate(geom, xoff=2_000_000, yoff=1_000_000)
        )
        return gdf.set_crs(TARGET_CRS, allow_override=True)
    raise ValueError(f"{path} could not be transformed to {TARGET_CRS}")


def _points_gdf(manifest: pd.DataFrame) -> gpd.GeoDataFrame:
    return gpd.GeoDataFrame(
        manifest[["global_index"]].copy(),
        geometry=gpd.points_from_xy(manifest["lv95_easting"], manifest["lv95_northing"]),
        crs=TARGET_CRS,
    )


def build_context(manifest: pd.DataFrame, spec: Mapping) -> JoinContext:
    """Bootstrap the lookups needed by every join_method."""
    ctx = JoinContext(manifest=manifest.copy(), biogeoregion_keys=spec.get("biogeoregion_keys", {}))

    # Municipality polygon: use LABES 36 Zersiedlung_2018_Gemeinde (already vetted as
    # canonical CH-wide municipality polygon with BFS-Nr; 2,222 features).
    muni_path = PROJECT_ROOT / "geodata/LABES/labes_36_zersiedlung/LABES_36_GIS/LABES_36_Zersiedlung_2018.gdb"
    muni_layer = "Zersiedlung_2018_Gemeinde"
    logger.info("[context] loading municipality polygon %s/%s", muni_path.name, muni_layer)
    munis = _read_layer(muni_path, layer=muni_layer)
    munis = munis[["bfs_nummer", "geometry"]].copy()
    munis["bfs_nummer"] = pd.to_numeric(munis["bfs_nummer"], errors="coerce").astype("Int64")

    biogreg_path = muni_path
    biogreg_layer = "biogeoregio_2018"
    logger.info("[context] loading biogeoregion polygon %s/%s", biogreg_path.name, biogreg_layer)
    biogregs = _read_layer(biogreg_path, layer=biogreg_layer)
    biogregs = biogregs[["biogreg_c6", "geometry"]].copy()
    biogregs["biogreg_c6"] = pd.to_numeric(biogregs["biogreg_c6"], errors="coerce").astype("Int64")

    points = _points_gdf(manifest)
    logger.info("[context] sjoin manifest -> municipality (%s pts × %s polys)", f"{len(points):,}", f"{len(munis):,}")
    muni_join = gpd.sjoin(points, munis, how="left", predicate="within")
    muni_join = muni_join[~muni_join.index.duplicated(keep="first")]
    ctx.municipality_bfs_per_point = (
        pd.Series(muni_join["bfs_nummer"].values, index=muni_join["global_index"].values, name="bfs_nummer")
    )

    logger.info("[context] sjoin manifest -> biogeoregion")
    bg_join = gpd.sjoin(points, biogregs, how="left", predicate="within")
    bg_join = bg_join[~bg_join.index.duplicated(keep="first")]
    ctx.biogeoregion_per_point = (
        pd.Series(bg_join["biogreg_c6"].values, index=bg_join["global_index"].values, name="biogreg_c6")
    )

    miss_m = int(ctx.municipality_bfs_per_point.isna().sum())
    miss_b = int(ctx.biogeoregion_per_point.isna().sum())
    logger.info(
        "[context] municipality lookup: %s/%s missing (%.2f%%)",
        f"{miss_m:,}", f"{len(manifest):,}", 100.0 * miss_m / len(manifest),
    )
    logger.info(
        "[context] biogeoregion lookup: %s/%s missing (%.2f%%)",
        f"{miss_b:,}", f"{len(manifest):,}", 100.0 * miss_b / len(manifest),
    )
    return ctx


# -------------------------------------------------------------------
# join_method dispatch
# -------------------------------------------------------------------


def social_csv_municipality(ctx: JoinContext, spec: Mapping) -> pd.Series:
    src = spec["source"]
    path = PROJECT_ROOT / src["path"]
    gemnr_col = src["gemnr_column"]
    if "value_column" in src and src["value_column"]:
        value_columns = [src["value_column"]]
        aggregation = src.get("aggregation", "municipality_mean")
    else:
        value_columns = src.get("value_columns") or []
        aggregation = src.get("aggregation", "municipality_mean_then_average")
    coverage_fallback = src.get("coverage_fallback")

    df = pd.read_csv(path, sep=None, engine="python", encoding="utf-8-sig")
    missing = [c for c in value_columns if c not in df.columns]
    if missing:
        raise KeyError(f"social CSV {path.name} missing columns {missing}")

    df[gemnr_col] = pd.to_numeric(df[gemnr_col], errors="coerce").astype("Int64")
    df = df.dropna(subset=[gemnr_col])
    for col in value_columns:
        df[col] = pd.to_numeric(
            df[col].astype(str).str.replace(",", ".", regex=False),
            errors="coerce",
        )

    if aggregation == "municipality_mean":
        per_respondent = df[value_columns[0]]
    elif aggregation == "municipality_mean_then_average":
        # average each respondent's value across the columns first, then per-municipality
        # mean of those individual averages.
        per_respondent = df[value_columns].mean(axis=1, skipna=True)
    else:
        raise ValueError(f"unknown aggregation {aggregation}")
    df = df.assign(_indiv=per_respondent)
    muni_table = df.groupby(gemnr_col, dropna=True)["_indiv"].mean()
    muni_table.index = muni_table.index.astype("Int64")

    bfs_per_point = ctx.municipality_bfs_per_point.reindex(ctx.manifest["global_index"].values)
    values = bfs_per_point.map(muni_table.to_dict())

    if coverage_fallback == "biogeoregion_mean":
        # Build per-biogeoregion mean of respondent-level values.
        # Need a Gemnr → biogeoregion mapping; derive it from the manifest's per-point
        # municipality+biogeoregion mapping by majority vote.
        muni_to_biogreg = (
            pd.DataFrame(
                {
                    "bfs": ctx.municipality_bfs_per_point.values,
                    "bg": ctx.biogeoregion_per_point.values,
                }
            )
            .dropna()
            .groupby("bfs")["bg"]
            .agg(lambda s: s.value_counts().idxmax())
        )
        df = df.assign(_bg=df[gemnr_col].map(muni_to_biogreg))
        biogreg_table = df.dropna(subset=["_bg"]).groupby("_bg")["_indiv"].mean()
        biogreg_per_point = ctx.biogeoregion_per_point.reindex(ctx.manifest["global_index"].values)
        fallback = biogreg_per_point.map(biogreg_table.to_dict())
        values = pd.Series(values.values).where(pd.Series(values.values).notna(), fallback.values)

    return pd.Series(values.values, index=ctx.manifest["global_index"].values, name=spec["column_name"])


def polygon_municipality_field(ctx: JoinContext, spec: Mapping) -> pd.Series:
    src = spec["source"]
    path = PROJECT_ROOT / src["path"]
    layer = src.get("layer")
    bfs_col = src["bfs_column"]
    value_col = src["value_column"]
    gdf = _read_layer(path, layer=layer)
    gdf[bfs_col] = pd.to_numeric(gdf[bfs_col], errors="coerce").astype("Int64")
    gdf[value_col] = pd.to_numeric(gdf[value_col], errors="coerce")
    table = gdf.dropna(subset=[bfs_col]).set_index(bfs_col)[value_col]
    table = table[~table.index.duplicated(keep="first")]
    bfs_per_point = ctx.municipality_bfs_per_point.reindex(ctx.manifest["global_index"].values)
    values = bfs_per_point.map(table.to_dict())
    return pd.Series(values.values, index=ctx.manifest["global_index"].values, name=spec["column_name"])


def polygon_biogeoregion_field(ctx: JoinContext, spec: Mapping) -> pd.Series:
    src = spec["source"]
    path = PROJECT_ROOT / src["path"]
    layer = src.get("layer")
    biogreg_col = src.get("biogreg_column", "biogreg_c6")
    value_col = src["value_column"]
    gdf = _read_layer(path, layer=layer)
    gdf[biogreg_col] = pd.to_numeric(gdf[biogreg_col], errors="coerce").astype("Int64")
    gdf[value_col] = pd.to_numeric(gdf[value_col], errors="coerce")
    table = gdf.dropna(subset=[biogreg_col]).set_index(biogreg_col)[value_col]
    table = table[~table.index.duplicated(keep="first")]
    biogreg_per_point = ctx.biogeoregion_per_point.reindex(ctx.manifest["global_index"].values)
    values = biogreg_per_point.map(table.to_dict())
    return pd.Series(values.values, index=ctx.manifest["global_index"].values, name=spec["column_name"])


def polygon_cell_field(ctx: JoinContext, spec: Mapping) -> pd.Series:
    """Point-in-polygon over a fine cell coverage (not a municipality boundary)."""
    src = spec["source"]
    path = PROJECT_ROOT / src["path"]
    layer = src.get("layer")
    value_col = src.get("value_column")
    gdf = _read_layer(path, layer=layer)
    if value_col is None:
        # auto-pick the first numeric non-geometry column
        candidates = [
            c for c in gdf.columns
            if c != "geometry" and pd.api.types.is_numeric_dtype(gdf[c])
        ]
        if not candidates:
            raise ValueError(f"polygon_cell_field needs value_column for {path}")
        value_col = candidates[0]
        logger.info("[%s] auto-selected value_column=%s", spec["column_name"], value_col)
    gdf[value_col] = pd.to_numeric(gdf[value_col], errors="coerce")
    keep = gdf[[value_col, "geometry"]].copy()
    points = _points_gdf(ctx.manifest)
    joined = gpd.sjoin(points, keep, how="left", predicate="within")
    joined = joined[~joined.index.duplicated(keep="first")]
    return pd.Series(
        joined[value_col].values,
        index=joined["global_index"].values,
        name=spec["column_name"],
    )


def raster_point_sample(ctx: JoinContext, spec: Mapping) -> pd.Series:
    src = spec["source"]
    path = PROJECT_ROOT / src["path"]
    scale = float(src.get("value_scale", 1.0))
    offset = float(src.get("value_offset", 0.0))
    nodata_sentinel_threshold = float(src.get("nodata_above", 1e30))
    # Optional class-membership mode: emit 1 if pixel value ∈ class_set else 0
    # (NaN preserved). Used to derive binary indicators from categorical land-use
    # rasters (e.g., ind3 settlement from NOAS04 codes {11, 12, 13, 15, 16, 17}).
    class_set = src.get("class_set")

    with rasterio.open(path) as ds:
        src_crs = ds.crs
        if src_crs is None:
            assumed = src.get("crs_assumed")
            if not assumed:
                raise ValueError(f"raster {path} has no CRS and no crs_assumed in spec")
            src_crs = rasterio.crs.CRS.from_string(assumed)
        manifest = ctx.manifest
        if str(src_crs) == TARGET_CRS or (hasattr(src_crs, "to_epsg") and src_crs.to_epsg() == 2056):
            xs = manifest["lv95_easting"].to_numpy()
            ys = manifest["lv95_northing"].to_numpy()
        else:
            xs, ys = rio_transform(
                rasterio.crs.CRS.from_string(TARGET_CRS),
                src_crs,
                manifest["lv95_easting"].to_numpy(),
                manifest["lv95_northing"].to_numpy(),
            )
        # masked=False because rasterio otherwise masks 0-valued cells as
        # nodata for Int8 binary mask rasters even when ds.nodata is None.
        sample = list(ds.sample(zip(xs, ys), masked=False))
        nodata = ds.nodata

    values = np.array([s[0] if s is not None else np.nan for s in sample], dtype=np.float64)
    # Treat file-nodata + Esri float sentinels (~3.4e38) as missing.
    values = np.where(np.abs(values) > nodata_sentinel_threshold, np.nan, values)
    if nodata is not None:
        values = np.where(values == nodata, np.nan, values)
    if class_set is not None:
        allowed = set(int(c) for c in class_set)
        mask = ~np.isnan(values)
        out = np.full_like(values, np.nan, dtype=np.float64)
        ints = np.where(mask, values.astype(np.int64), -1)
        out[mask] = np.array([1.0 if int(v) in allowed else 0.0 for v in ints[mask]])
        values = out
    if scale != 1.0 or offset != 0.0:
        values = values * scale + offset
    return pd.Series(values, index=manifest["global_index"].values, name=spec["column_name"])


def arealstatistik_point_sample(ctx: JoinContext, spec: Mapping) -> pd.Series:
    """Sample BFS Arealstatistik point data (100m lattice) at manifest points.

    The official dataset is a point lattice with coordinates on a 100m grid.
    We snap each manifest LV95 coordinate to the nearest 100m lattice point,
    then classify class-membership from the configured class column.
    """
    src = spec["source"]
    path = PROJECT_ROOT / src["path"]
    e_col = src.get("easting_column", "E_COORD")
    n_col = src.get("northing_column", "N_COORD")
    class_col = src.get("class_column", "AS_72")
    class_set = src.get("class_set")
    if not class_set:
        raise ValueError(f"{spec['column_name']} requires source.class_set")
    allowed = set(int(c) for c in class_set)
    grid_size_m = int(src.get("grid_size_m", 100))
    scale = float(src.get("value_scale", 1.0))
    offset = float(src.get("value_offset", 0.0))

    table = pd.read_csv(path, sep=";", usecols=[e_col, n_col, class_col], compression="infer")
    table[e_col] = pd.to_numeric(table[e_col], errors="coerce").astype("Int64")
    table[n_col] = pd.to_numeric(table[n_col], errors="coerce").astype("Int64")
    table[class_col] = pd.to_numeric(table[class_col], errors="coerce").astype("Int64")
    table = table.dropna(subset=[e_col, n_col, class_col])

    # Packed integer key for fast vectorized lookup.
    table_key = (
        table[e_col].astype(np.int64).to_numpy() * 10_000_000
        + table[n_col].astype(np.int64).to_numpy()
    )
    table_val = table[class_col].astype(np.int64).isin(allowed).to_numpy(dtype=np.float64)
    lookup = pd.Series(table_val, index=table_key)
    lookup = lookup[~lookup.index.duplicated(keep="first")]

    manifest = ctx.manifest
    ex = manifest["lv95_easting"].to_numpy(dtype=np.float64)
    ny = manifest["lv95_northing"].to_numpy(dtype=np.float64)
    ex_snap = (np.rint(ex / grid_size_m) * grid_size_m).astype(np.int64)
    ny_snap = (np.rint(ny / grid_size_m) * grid_size_m).astype(np.int64)
    manifest_key = ex_snap * 10_000_000 + ny_snap

    values = lookup.reindex(manifest_key).to_numpy(dtype=np.float64)
    if scale != 1.0 or offset != 0.0:
        values = values * scale + offset
    return pd.Series(values, index=manifest["global_index"].values, name=spec["column_name"])


def xlsx_biogeoregion_via_polygon(ctx: JoinContext, spec: Mapping) -> pd.Series:
    """Read an XLSX with one row per biogeoregion, map to manifest via biogreg code.

    Handles three BAFU XLSX patterns:
      1. Column with integer `biogreg_c6` codes (1..6).
      2. Column with biogeoregion names as text (Jura, Mittelland, …).
      3. Multi-header sheets (e.g. row 0 = section heading, row 1 = year, row 2+ = data)
         — pass `header_row: 2` and an optional `teilraum_column` in the spec.
    """
    src = spec["source"]
    path = PROJECT_ROOT / src["path"]
    sheet = src.get("sheet")
    header_row = src.get("header_row", 0)
    table = pd.read_excel(path, sheet_name=sheet, header=header_row)
    # collapse stray nan column names that confuse autodetection
    table.columns = [str(c) for c in table.columns]

    biogreg_col = (
        src.get("teilraum_column")
        or src.get("biogreg_column")
        or _autodetect_biogreg_column(table)
    )
    if biogreg_col is None:
        raise ValueError(f"xlsx {path.name} sheet={sheet}: no biogeoregion column found")
    biogreg_col = _resolve_column(table.columns, biogreg_col)
    if biogreg_col not in table.columns:
        raise ValueError(f"xlsx {path.name} sheet={sheet}: biogeoregion column {biogreg_col!r} not found")

    # Build code column from whichever form the source uses. Some BAFU sheets
    # publish a combined "Zentralalpen" row; duplicate it to both central-Alpine
    # bioregions because the manifest uses the six-way biogeoregion key.
    if pd.api.types.is_numeric_dtype(table[biogreg_col]):
        codes = pd.to_numeric(table[biogreg_col], errors="coerce").astype("Int64")
        table = table.assign(_bg_code=codes).dropna(subset=["_bg_code"])
    else:
        expanded_rows = []
        for _, row in table.iterrows():
            codes = _normalize_biogeoregion_text_set(row[biogreg_col])
            for code in codes:
                expanded = row.copy()
                expanded["_bg_code"] = code
                expanded_rows.append(expanded)
        table = pd.DataFrame(expanded_rows)
        if table.empty:
            raise ValueError(f"xlsx {path.name} sheet={sheet}: no biogeoregion rows found")

    value_col = src.get("value_column") or _autodetect_numeric_value_column(
        table, exclude={biogreg_col, "_bg_code"}
    )
    if value_col is not None:
        value_col = _resolve_column(table.columns, value_col)
    if value_col is None:
        raise ValueError(f"xlsx {path.name} sheet={sheet}: no numeric value column found")
    if value_col not in table.columns:
        raise ValueError(f"xlsx {path.name} sheet={sheet}: value column {value_col!r} not found")
    table[value_col] = _coerce_numeric(table[value_col])

    mapping = table.groupby("_bg_code", dropna=True)[value_col].mean().to_dict()
    scale = float(src.get("value_scale", 1.0))
    if scale != 1.0:
        mapping = {k: float(v) * scale for k, v in mapping.items() if pd.notna(v)}
    logger.info(
        "[%s] xlsx biogeoregion mapping (scale=%g): %s",
        spec["column_name"], scale,
        {int(k): float(v) for k, v in mapping.items() if pd.notna(v)},
    )

    biogreg_per_point = ctx.biogeoregion_per_point.reindex(ctx.manifest["global_index"].values)
    values = biogreg_per_point.map(mapping)
    return pd.Series(values.values, index=ctx.manifest["global_index"].values, name=spec["column_name"])


def sample_based_biogeoregion(ctx: JoinContext, spec: Mapping) -> pd.Series:
    """Alias — same dispatch as xlsx_biogeoregion_via_polygon."""
    return xlsx_biogeoregion_via_polygon(ctx, spec)


def xlsx_biogeoregion_rows(ctx: JoinContext, spec: Mapping) -> pd.Series:
    """Read one or more labelled rows from a matrix-style XLSX table.

    Some LABES workbooks put biogeoregions across columns and indicators down
    rows, often with repeated unit/helper columns. The spec declares which row
    contains region names and which row labels to aggregate.
    """
    src = spec["source"]
    path = PROJECT_ROOT / src["path"]
    sheet = src.get("sheet")
    region_row = int(src["region_row"])
    row_label_column = int(src.get("row_label_column", 0))
    labels = [str(v).strip() for v in (src.get("row_labels") or [])]
    if not labels:
        raise ValueError(f"xlsx_biogeoregion_rows needs row_labels for {spec['column_name']}")
    aggregate = src.get("aggregate", "sum")
    unit_row = src.get("unit_row")
    unit_marker = src.get("value_unit_marker")

    table = pd.read_excel(path, sheet_name=sheet, header=None)
    region_cols: dict[int, int] = {}
    for col in range(table.shape[1]):
        code = _normalize_biogeoregion_text(table.iat[region_row, col])
        if code is None:
            continue
        if unit_row is not None and unit_marker is not None:
            unit = str(table.iat[int(unit_row), col]).strip()
            if unit != str(unit_marker):
                continue
        region_cols[col] = code
    if not region_cols:
        raise ValueError(f"{path.name} sheet={sheet}: no biogeoregion columns found")

    row_values = []
    label_series = table.iloc[:, row_label_column].astype(str).str.strip()
    for label in labels:
        matches = table[label_series == label]
        if matches.empty:
            raise ValueError(f"{path.name} sheet={sheet}: row label {label!r} not found")
        row_values.append(matches.iloc[0])

    mapping: dict[int, float] = {}
    for col, code in region_cols.items():
        values = [_coerce_numeric(pd.Series([row[col]])).iloc[0] for row in row_values]
        values = [v for v in values if pd.notna(v)]
        if not values:
            continue
        if aggregate == "mean":
            mapping[code] = float(np.mean(values))
        elif aggregate == "first":
            mapping[code] = float(values[0])
        else:
            mapping[code] = float(np.sum(values))

    logger.info("[%s] xlsx row biogeoregion mapping: %s", spec["column_name"], mapping)
    biogreg_per_point = ctx.biogeoregion_per_point.reindex(ctx.manifest["global_index"].values)
    values = biogreg_per_point.map(mapping)
    return pd.Series(values.values, index=ctx.manifest["global_index"].values, name=spec["column_name"])


def xlsx_national_constant(ctx: JoinContext, spec: Mapping) -> pd.Series:
    """Single national scalar; emit as a constant column."""
    ref = (spec.get("reference_values") or {}).get("national_mean")
    if ref is None:
        # try to read from XLSX file as a last resort
        src = spec["source"]
        path = PROJECT_ROOT / src["path"]
        sheet = src.get("sheet")
        try:
            if src.get("cell"):
                from openpyxl import load_workbook
                from openpyxl.utils.cell import coordinate_to_tuple

                wb = load_workbook(path, read_only=True, data_only=True)
                ws = wb[sheet] if sheet else wb.active
                row, col = coordinate_to_tuple(str(src["cell"]))
                ref = _coerce_numeric(pd.Series([ws.cell(row=row, column=col).value])).iloc[0]
            else:
                t = pd.read_excel(path, sheet_name=sheet)
                numeric_cols = t.select_dtypes(include="number").columns
                for col in numeric_cols:
                    vals = _coerce_numeric(t[col]).dropna()
                    if len(vals):
                        ref = float(vals.iloc[-1])
                        break
        except Exception as exc:
            logger.warning("[%s] xlsx_national_constant read failed: %s", spec["column_name"], exc)
    return pd.Series(
        np.full(len(ctx.manifest), ref if ref is not None else np.nan, dtype=np.float64),
        index=ctx.manifest["global_index"].values,
        name=spec["column_name"],
    )


def spec_biogeoregion_values(ctx: JoinContext, spec: Mapping) -> pd.Series:
    """Map PDF/reference values embedded in the spec to manifest biogeoregions."""
    refs = (spec.get("reference_values") or {}).get("biogeoregion_means") or {}
    mapping = {}
    for name, value in refs.items():
        code = ctx.biogeoregion_keys.get(name) or _normalize_biogeoregion_text(name)
        if code is not None and value is not None:
            mapping[int(code)] = float(value)
    if not mapping:
        raise ValueError(f"{spec['column_name']} has no biogeoregion reference values")
    biogreg_per_point = ctx.biogeoregion_per_point.reindex(ctx.manifest["global_index"].values)
    values = biogreg_per_point.map(mapping)
    return pd.Series(values.values, index=ctx.manifest["global_index"].values, name=spec["column_name"])


def nearest_feature_distance(ctx: JoinContext, spec: Mapping) -> pd.Series:
    """Distance in metres from each manifest point to the nearest source feature."""
    src = spec["source"]
    path = PROJECT_ROOT / src["path"]
    layer = src.get("layer")
    features = _read_layer(path, layer=layer)
    filter_col = src.get("filter_column")
    filter_values = src.get("filter_values")
    if filter_col and filter_values is not None:
        features = features[features[filter_col].isin(filter_values)].copy()
    features = features[~features.geometry.isna() & ~features.geometry.is_empty]
    if features.empty:
        raise ValueError(f"nearest_feature_distance source is empty for {spec['column_name']}")

    points = _points_gdf(ctx.manifest)
    try:
        joined = gpd.sjoin_nearest(
            points,
            features[["geometry"]],
            how="left",
            distance_col="_distance_m",
        )
        joined = joined[~joined.index.duplicated(keep="first")]
        values = joined.set_index("global_index")["_distance_m"].reindex(ctx.manifest["global_index"].values)
    except Exception as exc:
        logger.warning("[%s] sjoin_nearest failed, falling back to union distance: %s", spec["column_name"], exc)
        union = features.geometry.union_all() if hasattr(features.geometry, "union_all") else features.geometry.unary_union
        values = points.geometry.distance(union)
    return pd.Series(values.values, index=ctx.manifest["global_index"].values, name=spec["column_name"])


# -------------------------------------------------------------------
# Helpers
# -------------------------------------------------------------------


_BIOGREG_NAME_RE = re.compile(r"biogreg|biogeo|teilraum|region|grossregion", re.I)

# Map every plausible BAFU spelling of a biogeoregion → numeric code 1..6.
BIOGEOREGION_NAME_TO_CODE = {}
for code, canonical in BIOGEOREGION_NAMES.items():
    BIOGEOREGION_NAME_TO_CODE[canonical] = code
_BG_ALIASES = {
    1: ["jura"],
    2: ["mittelland", "plateau", "central plateau", "mittellaender"],
    3: ["alpennordflanke", "alpennord", "voralpen", "alpen nord"],
    4: ["westliche zentralalpen", "westl zentralalpen", "westl. zentralalpen", "westl z", "westl. z.", "westliche z", "westliche z.", "zentralalpen west"],
    5: ["oestliche zentralalpen", "oestl zentralalpen", "oestl. zentralalpen", "oestl z", "oestl. z.", "oestliche z", "östliche zentralalpen", "östl zentralalpen", "östl. zentralalpen", "östl z", "östl. z.", "östliche z", "ostliche zentralalpen", "ostliche z", "ost. z.", "ost zentralalpen", "zentralalpen ost"],
    6: ["alpensuedflanke", "alpensüdflanke", "alpensüd", "südseite", "suedseite"],
}
for code, aliases in _BG_ALIASES.items():
    for a in aliases:
        BIOGEOREGION_NAME_TO_CODE[a] = code


def _normalize_biogeoregion_text(text) -> Optional[int]:
    """Map a possibly-fuzzy biogeoregion label to its 1..6 code."""
    if pd.isna(text):
        return None
    s = str(text).strip().lower()
    s = s.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    s = re.sub(r"\s+", " ", s)
    s = s.rstrip(".")
    if s in BIOGEOREGION_NAME_TO_CODE:
        return BIOGEOREGION_NAME_TO_CODE[s]
    # try fuzzy contains check (e.g., "Östliche Zentralalpen 2020" → 5)
    for alias, code in BIOGEOREGION_NAME_TO_CODE.items():
        if alias in s:
            return code
    return None


def _normalize_biogeoregion_text_set(text) -> list[int]:
    code = _normalize_biogeoregion_text(text)
    if code is not None:
        return [code]
    if pd.isna(text):
        return []
    s = str(text).strip().lower()
    s = s.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    s = re.sub(r"\s+", " ", s).rstrip(".")
    if s in {"zentralalpen", "zentrale alpen", "zentral alpen"}:
        return [4, 5]
    return []


def _autodetect_biogreg_column(table: pd.DataFrame) -> Optional[str]:
    # First pass: column name match
    for c in table.columns:
        if _BIOGREG_NAME_RE.search(str(c)):
            return c
    # Second pass: any column whose values include canonical biogeoregion names
    for c in table.columns:
        if not table[c].dtype == object:
            continue
        hits = sum(_normalize_biogeoregion_text(v) is not None for v in table[c].dropna().head(20))
        if hits >= 3:
            return c
    return None


def _normalize_col_label(value: Any) -> str:
    s = str(value).strip().lower().replace("*", "")
    return re.sub(r"\s+", " ", s)


def _resolve_column(columns, requested) -> str:
    """Resolve YAML column labels against pandas' sometimes-munged XLSX labels."""
    if requested in columns:
        return requested
    req_norm = _normalize_col_label(requested)
    for col in columns:
        if _normalize_col_label(col) == req_norm:
            return col
    for col in columns:
        if req_norm and req_norm in _normalize_col_label(col):
            return col
    return str(requested)


def _coerce_numeric(series: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce")
    return pd.to_numeric(
        series.astype(str)
        .str.replace("'", "", regex=False)
        .str.replace(",", ".", regex=False)
        .str.replace("%", "", regex=False),
        errors="coerce",
    )


def _autodetect_numeric_value_column(table: pd.DataFrame, exclude=()) -> Optional[str]:
    for c in table.columns:
        if c in exclude:
            continue
        if pd.api.types.is_numeric_dtype(table[c]):
            return c
    return None


NEU17_ELEMENTS = {
    "f30001": "baeche_fluesse",
    "f30002": "seen_weiher",
    "f30003": "paerke_gruenanlagen",
    "f30004": "wiesen_aecker_felder",
    "f30005": "wald",
    "f30006": "industrie_gewerbe",
    "f30007": "rebberge",
    "f30008": "historischer_dorfkern",
    "f30009": "berge_taeler",
    "f30010": "hochspannungsmasten",
    "f30011": "windenergie",
    "f30012": "solarenergie",
    "f30013": "moore",
    "f30014": "bauernhoefe_staelle",
    "f30015": "burg_schloss",
    "f30016": "streusiedlungen",
    "f30017": "naturschutzgebiet",
    "f30018": "flugplatz",
    "f30019": "hochhaeuser",
    "f30020": "wohnquartiere",
}

NEU18_CHANGES = {
    "f70001": "wohngebiete_neu",
    "f70002": "industrie_gewerbe_neu",
    "f70003": "strassen_neu",
    "f70004": "wege_asphaltiert",
    "f70005": "solaranlagen_neu",
    "f70006": "windkraftwerke_neu",
    "f70007": "wasserkraftwerke_neu",
    "f70008": "siedlung_dichter",
    "f70009": "gruenflaechen_verloren",
    "f70010": "siedlung_ausgebreitet",
    "f70011": "fliessgewaesser_renaturiert",
    "f70012": "fliessgewaesser_eingedolt",
    "f70013": "gewaesser_zugaenglicher",
    "f70014": "lw_eintoeniger",
    "f70015": "lw_vielfaeltiger",
    "f70016": "waldflaeche_zugenommen",
    "f70017": "traditionelle_strukturen_verloren",
    "f70018": "weniger_sterne_sichtbar",
    "f70019": "naturnahe_erholung_zugaenglicher",
}


def social_csv_vector_to_string(ctx: JoinContext, spec: Mapping) -> pd.Series:
    """Compress a multi-element survey vector to a pipe-delimited string per
    municipality.

    Used for NEU_17 Landschaftscharakter (20 yes/no items) and NEU_18
    Landschaftsveränderung (19 yes/no change-noticed items). BAFU does not
    publish these as scalar indices; we emit a per-municipality string of the
    element slugs that >threshold of respondents marked, plus a biogeoregion
    fallback when a municipality has no respondents.
    """
    src = spec["source"]
    path = PROJECT_ROOT / src["path"]
    gemnr_col = src["gemnr_column"]
    element_map: dict = src["element_map"]
    threshold = float(src.get("majority_threshold", 0.5))

    df = pd.read_csv(path, sep=None, engine="python", encoding="utf-8-sig")
    missing = [c for c in element_map if c not in df.columns]
    if missing:
        raise KeyError(f"social CSV {path.name} missing element columns {missing}")
    df[gemnr_col] = pd.to_numeric(df[gemnr_col], errors="coerce").astype("Int64")
    df = df.dropna(subset=[gemnr_col])

    # Binarise each element column. NEU_17 is 0=Nein, 1=Ja, so "yes_1" matches.
    # NEU_18 is 1-5 Likert + 6=KeineAntwort; "likert_ge4" means change noticed
    # at least with "trifft eher zu".
    rule = src.get("binarize_rule", "yes_1")
    def _binarize(s: pd.Series) -> pd.Series:
        v = pd.to_numeric(s, errors="coerce")
        if rule == "yes_1":
            return v == 1
        if rule == "likert_ge4":
            return (v >= 4) & (v <= 5)
        if rule == "likert_ge3":
            return (v >= 3) & (v <= 5)
        raise ValueError(f"unknown binarize_rule {rule!r}")

    bin_df = pd.DataFrame({slug: _binarize(df[code]) for code, slug in element_map.items()})
    bin_df["Gemnr"] = df[gemnr_col].values

    grouped = bin_df.groupby("Gemnr").mean()  # share of respondents per element
    sorted_slugs = list(element_map.values())

    def encode_row(row: pd.Series) -> str:
        active = [s for s in sorted_slugs if row.get(s, 0) >= threshold]
        return "|".join(active) if active else ""

    muni_table = grouped.apply(encode_row, axis=1)
    muni_table.index = muni_table.index.astype("Int64")

    # Biogeoregion fallback: encode the biogeoregion-level majority instead.
    muni_to_biogreg = (
        pd.DataFrame({"bfs": ctx.municipality_bfs_per_point.values, "bg": ctx.biogeoregion_per_point.values})
        .dropna()
        .groupby("bfs")["bg"]
        .agg(lambda s: s.value_counts().idxmax())
    )
    bin_df = bin_df.assign(_bg=bin_df["Gemnr"].map(muni_to_biogreg))
    bg_grouped = bin_df.dropna(subset=["_bg"]).groupby("_bg").mean(numeric_only=True)
    bg_table = bg_grouped.apply(encode_row, axis=1)

    bfs_per_point = ctx.municipality_bfs_per_point.reindex(ctx.manifest["global_index"].values)
    primary = bfs_per_point.map(muni_table.to_dict())
    biogreg_per_point = ctx.biogeoregion_per_point.reindex(ctx.manifest["global_index"].values)
    fallback = biogreg_per_point.map(bg_table.to_dict())
    final = pd.Series(primary.values).where(pd.Series(primary.values).notna() & (pd.Series(primary.values) != ""), fallback.values)
    return pd.Series(final.values, index=ctx.manifest["global_index"].values, name=spec["column_name"])


def _unsupported(ctx: JoinContext, spec: Mapping) -> pd.Series:
    """Indicators BAFU does not publish as a single scalar (vector parameters,
    per-statement parameters). Emit all-null with provenance in logs."""
    logger.info(
        "[%s] join_method=%s — emitting null column (see spec english_description)",
        spec["column_name"], spec.get("join_method"),
    )
    return pd.Series(
        np.full(len(ctx.manifest), np.nan, dtype=np.float64),
        index=ctx.manifest["global_index"].values,
        name=spec["column_name"],
    )


JOIN_DISPATCH = {
    "social_csv_municipality": social_csv_municipality,
    "polygon_municipality_field": polygon_municipality_field,
    "polygon_biogeoregion_field": polygon_biogeoregion_field,
    "polygon_cell_field": polygon_cell_field,
    "raster_point_sample": raster_point_sample,
    "arealstatistik_point_sample": arealstatistik_point_sample,
    "xlsx_biogeoregion_via_polygon": xlsx_biogeoregion_via_polygon,
    "xlsx_biogeoregion_rows": xlsx_biogeoregion_rows,
    "sample_based_biogeoregion": sample_based_biogeoregion,
    "xlsx_national_constant": xlsx_national_constant,
    "spec_biogeoregion_values": spec_biogeoregion_values,
    "nearest_feature_distance": nearest_feature_distance,
    "social_csv_vector_to_string": social_csv_vector_to_string,
    "unsupported_vector_parameter": _unsupported,
    "unsupported_no_single_index": _unsupported,
}


def dispatch(ctx: JoinContext, spec: Mapping) -> pd.Series:
    method = spec.get("join_method")
    fn = JOIN_DISPATCH.get(method)
    if fn is None:
        raise ValueError(f"unknown join_method {method!r} for {spec.get('column_name')}")
    t0 = time.time()
    series = fn(ctx, spec)
    dt = time.time() - t0
    logger.info(
        "[%s] joined via %s in %.1fs (%.2f%% non-null)",
        spec["column_name"], method, dt, 100.0 * series.notna().mean(),
    )
    return series


def sample_landmarks(
    spec: Mapping,
    landmarks: list[Mapping],
    indicator_value_by_global_index: pd.Series,
) -> dict[str, Optional[float]]:
    """Sample values at landmark coordinates where the join method supports it."""
    out: dict[str, Optional[float]] = {}
    if spec.get("join_method") == "raster_point_sample":
        src = spec["source"]
        path = PROJECT_ROOT / src["path"]
        try:
            with rasterio.open(path) as ds:
                lons = np.array([lm["lon"] for lm in landmarks], dtype=np.float64)
                lats = np.array([lm["lat"] for lm in landmarks], dtype=np.float64)
                src_crs = ds.crs or rasterio.crs.CRS.from_string(
                    src.get("crs_assumed", "EPSG:4326")
                )
                if str(src_crs) != "EPSG:4326":
                    xs, ys = rio_transform(
                        rasterio.crs.CRS.from_string("EPSG:4326"), src_crs, lons, lats
                    )
                else:
                    xs, ys = lons, lats
                samples = list(ds.sample(zip(xs, ys), masked=True))
                for lm, sample in zip(landmarks, samples):
                    if sample is None or (hasattr(sample, "mask") and bool(np.any(sample.mask))):
                        out[lm["id"]] = None
                    else:
                        out[lm["id"]] = float(sample[0])
        except Exception as exc:
            logger.warning("[%s] landmark raster sample failed: %s", spec.get("column_name"), exc)
        return out
    return out
